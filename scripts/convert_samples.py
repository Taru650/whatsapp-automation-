#!/usr/bin/env python3
"""Convert the district's print-layout Sonpur Mela workbook into the bot template.

  python3 scripts/convert_samples.py data/samples/Sonpur_Mela_Data_2025.xlsx

Writes
  data/templates/Sonpur_Mela_Bot_Data.xlsx  - upload to Google Drive, open as Google Sheets
  tests/fixtures/mela_sheet.json            - the same tabs as the Sheets API returns them
and prints what needs a human (legacy Kruti Dev text, missing values).

Structural clean-up only: merged cells are expanded, two-row headers and shift
column groups are flattened, multi-person cells are split into phone/phone_2.
Values are NOT corrected: a wrong phone stays wrong so the sync validation
(n8n/src/services/mela_sheet.js) flags it for the data owner.
"""
import datetime as dt
import json
import pathlib
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

PLACES_COLS = ["id", "category", "name_en", "name_hi", "location_en", "location_hi", "lat", "lon", "hours",
               "s1_name", "s1_phone", "s2_name", "s2_phone", "s3_name", "s3_phone", "allday_name", "allday_phone",
               "phone_2", "notes", "active", "verified", "updated_by"]
CONTROL_COLS = ["desk", "s1_name", "s1_designation", "s1_phone", "s2_name", "s2_designation", "s2_phone",
                "s3_name", "s3_designation", "s3_phone", "allday_name", "allday_designation", "allday_phone",
                "phone_2", "active", "verified", "updated_by"]
EVENT_COLS = ["id", "date", "time", "programme_en", "programme_hi", "artists", "department", "venue",
              "is_highlight", "active", "verified", "updated_by"]
GUIDE_COLS = ["id", "kind", "text_en", "text_hi", "sort", "active", "verified", "updated_by"]
KRUTI = re.compile(r"\s*¼[^½]*½")
notes = []


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip().rstrip("–-").strip()


def clean_name(sheet, row, v):
    text = s(v)
    if KRUTI.search(text):
        fixed = KRUTI.sub("", text).strip()
        notes.append(f"{sheet} row {row}: removed legacy Kruti Dev text from {text!r} -> {fixed!r}; "
                     "type the Hindi name in Unicode in name_hi")
        return fixed
    return text


def split_phones(v):
    parts = [p.strip() for p in re.split(r"&|,|/", s(v)) if p.strip()]
    return (parts + ["", ""])[:2]


def merged_value(ws, row, col):
    """Value of a cell, taking the top-left value of a merged range it belongs to."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value, rng.max_row > rng.min_row
    return ws.cell(row, col).value, False


def shift_places(ws, prefix, category, first_row):
    rows = []
    for r in range(first_row, ws.max_row + 1):
        name = clean_name(ws.title, r, ws.cell(r, 2).value)
        if not name:
            continue
        row = dict.fromkeys(PLACES_COLS, "")
        row.update(id=f"{prefix}{len(rows) + 1:02d}", category=category, name_en=name, active="TRUE", verified="FALSE")
        for i, sh in enumerate(("s1", "s2", "s3")):
            row[f"{sh}_name"] = s(ws.cell(r, 3 + 2 * i).value)
            row[f"{sh}_phone"] = s(ws.cell(r, 4 + 2 * i).value)
        rows.append(row)
    return rows


def simple_places(ws, prefix, category, first_row, name_col=2, person_col=None, phone_col=None):
    rows = []
    for r in range(first_row, ws.max_row + 1):
        name = clean_name(ws.title, r, ws.cell(r, name_col).value)
        if not name:
            continue
        row = dict.fromkeys(PLACES_COLS, "")
        row.update(id=f"{prefix}{len(rows) + 1:02d}", category=category, name_en=name, active="TRUE", verified="FALSE")
        if person_col:
            row["allday_name"] = s(ws.cell(r, person_col).value)
            row["allday_phone"] = s(ws.cell(r, phone_col).value)
        rows.append(row)
    return rows


def control_room(ws):
    desks = [("Magistrate", 2), ("Police", 4), ("Sanitation & Water", 6), ("Electricity", 8), ("Health", 10)]
    shift_rows = {1: 4, 2: 5, 3: 6}
    rows = []
    for desk, col in desks:
        row = dict.fromkeys(CONTROL_COLS, "")
        row.update(desk=desk, active="TRUE", verified="FALSE")
        name, merged = merged_value(ws, shift_rows[1], col)
        if merged:  # one merged cell over all three shifts = the same person all day
            phones, _ = merged_value(ws, shift_rows[1], col + 1)
            p1, p2 = split_phones(phones)
            row.update(allday_name=s(name), allday_phone=p1, phone_2=p2)
            notes.append(f"Control Room / {desk}: merged cell -> same officer(s) on every shift")
        else:
            for sh, r in shift_rows.items():
                p1, p2 = split_phones(ws.cell(r, col + 1).value)
                row[f"s{sh}_name"] = s(ws.cell(r, col).value)
                row[f"s{sh}_phone"] = p1
                if p2:
                    row["phone_2"] = p2
        rows.append(row)
    return rows


def events(ws):
    rows = []
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 2).value
        if not d:
            continue
        date = d.date().isoformat() if isinstance(d, dt.datetime) else s(d)
        row = dict.fromkeys(EVENT_COLS, "")
        row.update(id=f"E{len(rows) + 1:02d}", date=date, programme_en=s(ws.cell(r, 4).value),
                   department=s(ws.cell(r, 3).value), is_highlight="FALSE", active="TRUE", verified="FALSE")
        rows.append(row)
    return rows


def as_values(cols, rows):
    return [cols] + [[r.get(c, "") for c in cols] for r in rows]


def main(src):
    wb = load_workbook(src, data_only=True)
    places = (shift_places(wb["Thana"], "TH", "thana", 4)
              + shift_places(wb["Temporary Health Centre"], "HC", "health_centre", 4)
              + simple_places(wb["Veterinary Camp"], "VC", "vet_camp", 2, 2, 3, 4)
              + simple_places(wb["Parking"], "PK", "parking", 2)
              + simple_places(wb["Ghat"], "GH", "ghat", 3))
    tabs = {
        "places": as_values(PLACES_COLS, places),
        "control_room": as_values(CONTROL_COLS, control_room(wb["Control Room"])),
        "events": as_values(EVENT_COLS, events(wb["Mela Schedule"])),
        "guidelines": [GUIDE_COLS],
        "settings": [["key", "value"],
                     ["shift1_start", "06:00"], ["shift2_start", "14:00"], ["shift3_start", "22:00"],
                     ["public_helpline_1", ""], ["public_helpline_2", ""],
                     ["mela_center_lat", ""], ["mela_center_lon", ""], ["mela_radius_km", "5"],
                     ["mela_start", ""], ["mela_end", ""]],
    }
    notes.append("settings: public_helpline_1 is empty - the district must supply the public control room number")
    notes.append("settings: mela_center_lat/lon are empty - needed for the 'outside the Mela area' check")
    notes.append("all rows are marked verified=FALSE: the data owner ticks each row after checking it")

    fixture = ROOT / "tests" / "fixtures" / "mela_sheet.json"
    fixture.write_text(json.dumps(tabs, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    write_workbook(tabs, ROOT / "data" / "templates" / "Sonpur_Mela_Bot_Data.xlsx")
    counts = {k: len(v) - 1 for k, v in tabs.items()}
    print(f"converted: {counts}")
    print("needs a human:")
    for n in notes:
        print(" -", n)
    print(f"wrote {fixture.relative_to(ROOT)} and data/templates/Sonpur_Mela_Bot_Data.xlsx")
    print("next: node scripts/validate_mela.mjs tests/fixtures/mela_sheet.json   (shows every row the sync will reject)")


README_ROWS = [
    ("Sonpur Mela bot data / सोनपुर मेला बॉट डेटा", ""),
    ("", ""),
    ("1. Row 1 holds the headers. Never rename, merge or reorder them.", "पहली पंक्ति के शीर्षक कभी न बदलें, न मिलाएँ, न क्रम बदलें।"),
    ("2. One table per tab. No title rows, no blank rows in between.", "हर टैब में एक ही तालिका। बीच में खाली पंक्ति न छोड़ें।"),
    ("3. Phone: 10-digit mobile (9431012345) or landline with STD code and 0 (06152-240001). One number per cell; second number in phone_2.",
     "फ़ोन: 10 अंकों का मोबाइल या STD कोड के साथ लैंडलाइन (0 सहित)। एक सेल में एक ही नंबर; दूसरा phone_2 में।"),
    ("4. Dates as 2026-11-24 or 24/11/2026; times as 18:30.", "तारीख 2026-11-24 या 24/11/2026; समय 18:30।"),
    ("5. Hindi must be typed in Unicode (Google Input Tools / Windows Hindi keyboard), never Kruti Dev.",
     "हिंदी यूनिकोड में टाइप करें (Google Input Tools), Kruti Dev में नहीं।"),
    ("6. Use the dropdowns where offered.", "जहाँ ड्रॉपडाउन है, वहीं से चुनें।"),
    ("7. active = FALSE hides a row (don't delete rows).", "पंक्ति छिपाने के लिए active = FALSE करें, पंक्ति न मिटाएँ।"),
    ("8. Tick verified only after calling the number / checking the order. Unverified rows never reach citizens.",
     "नंबर पर फ़ोन करके या आदेश से मिलाकर ही verified = TRUE करें।"),
    ("9. Write your name in updated_by.", "updated_by में अपना नाम लिखें।"),
    ("10. id (TH01, HC03...) is permanent - never reuse it.", "id (TH01, HC03...) स्थायी है, दोबारा प्रयोग न करें।"),
    ("", ""),
    ("Shift columns: s1 = shift 1, s2 = shift 2, s3 = shift 3 (times in the settings tab). allday = same person all day.",
     "s1/s2/s3 = पहली/दूसरी/तीसरी पाली (समय settings टैब में)। allday = पूरे दिन वही व्यक्ति।"),
    ("Changes reach the bot within 10 minutes. If a row has an error, the bot keeps the old data and the admins get a WhatsApp alert with the row number.",
     "बदलाव 10 मिनट में बॉट तक पहुँचते हैं। गलती होने पर बॉट पुराना डेटा रखता है और एडमिन को पंक्ति नंबर के साथ संदेश जाता है।"),
]
PHONE_COLS = {"s1_phone", "s2_phone", "s3_phone", "allday_phone", "phone_2", "value"}
DROPDOWNS = {
    "category": '"thana,health_centre,vet_camp,parking,ghat,accommodation,toilet_water,lost_found,help_desk"',
    "desk": '"Magistrate,Police,Sanitation & Water,Electricity,Health,Other"',
    "kind": '"do,dont,emergency"',
    "active": '"TRUE,FALSE"', "verified": '"TRUE,FALSE"', "is_highlight": '"TRUE,FALSE"',
}


def write_workbook(tabs, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for row in README_ROWS:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 110
    ws.column_dimensions["B"].width = 90
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    head_fill = PatternFill("solid", fgColor="DDE6F7")
    for name, values in tabs.items():
        ws = wb.create_sheet(name)
        for row in values:
            ws.append(row)
        ws.freeze_panes = "A2"
        header = values[0]
        for j, col in enumerate(header, start=1):
            c = ws.cell(1, j)
            c.font = Font(bold=True)
            c.fill = head_fill
            letter = c.column_letter
            ws.column_dimensions[letter].width = 34 if col.endswith(("name", "_en", "_hi", "artists", "programme_en")) else 16
            if col in PHONE_COLS or col == "date":
                for r in range(2, 501):  # plain text: keeps leading zeros and exact digits
                    ws.cell(r, j).number_format = "@"
            if col in DROPDOWNS:
                dv = DataValidation(type="list", formula1=DROPDOWNS[col], allow_blank=True)
                dv.add(f"{letter}2:{letter}500")
                ws.add_data_validation(dv)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ROOT / "data" / "samples" / "Sonpur_Mela_Data_2025.xlsx")
